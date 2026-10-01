"""Render a KPIntradayReport as a Word document.

Mirrors the section order and prose of the reference "KP INTRADAY TREND CHANGE
PREDICTION REPORT" (Astro Quant Research Lab, v15.1) so the engine's output can
be shared exactly the way the hand-produced reports are.
"""

from __future__ import annotations

import io
from typing import Iterable

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

_GREY = RGBColor(0x55, 0x5B, 0x66)


def _title(doc, text, size=16, bold=True, align=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def _meta(doc, text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(9.5)
    r.font.color.rgb = _GREY
    return p


def _h(doc, text):
    return doc.add_heading(text, level=2)


def _h3(doc, text):
    return doc.add_heading(text, level=3)


def _para(doc, text, size=10.5):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(size)
    return p


def _table(doc, header: Iterable[str], rows: Iterable[Iterable[str]], size=8.5):
    header = list(header)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(str(h))
        run.bold = True
        run.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run("" if v is None else str(v))
            run.font.size = Pt(size)
    return t


def render_docx(report) -> bytes:
    ds = report.day_signature
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Pt(54)
        s.top_margin = s.bottom_margin = Pt(50)

    # --- title block ------------------------------------------------------
    _title(doc, "KP INTRADAY TREND CHANGE PREDICTION REPORT", 16)
    _meta(doc, f"Trading Day: {report.trading_day.isoformat()}")
    _meta(doc, f"{report.instrument}  •  {report.market}  •  {report.ayanamsa}")
    _meta(doc, f"Source: {report.source_file}   |   Generated: "
               f"{report.generated_at.strftime('%d-%b-%Y %H:%M')}")
    _meta(doc, report.brand)

    # --- 1. Day chart signature ------------------------------------------
    _h(doc, "1. Day Chart Signature")
    _para(doc, (f"Day Lord: {ds.day_lord}.   Moon: {ds.moon_sign} sign, {ds.moon_star} star "
                f"(sig strength {ds.day_bias_score:+.2f}), opening sub {ds.opening_sublord}.   "
                f"Moon speed: {ds.moon_speed:.2f} deg/day."))
    _para(doc, ds.day_bias_label)
    if getattr(ds, "rahu_windows_note", ""):
        _para(doc, ds.rahu_windows_note)
    _para(doc, ds.moon_affliction_gate)
    _para(doc, "Hora timeline with hora-lord aspects (v12.2, market hours):")
    for hs in ds.hora_slots:
        base = getattr(hs, "base_quality", hs.quality.value)
        interp = getattr(hs, "interpretation", hs.aspect_notes) or ""
        asp = getattr(hs, "aspects_received", "") or ""
        line = f"{hs.time_range_str()}  {hs.lord} ({base})."
        line += (f"  Lord receives: {asp} -> {interp}" if asp and asp != "None"
                 else "  Lord receives no major aspect — base hora quality applies unmodified.")
        _para(doc, line, 9.5)
    _para(doc, ("Rule under test: base hora quality is modified by aspects the hora lord receives — "
                "Mars aspect = trending/impulsive, Rahu = trap/false-move, Jupiter = support (can override "
                "a malefic base), Saturn = suppression. Hora contributes weight 0.3 to slot scores. "
                "Do not treat hora or its aspects as a standalone signal."), 9)
    if getattr(ds, "vargottama_note", ""):
        _para(doc, ds.vargottama_note)

    # --- 1b. Navamsa ------------------------------------------------------
    nav = getattr(ds, "navamsa", None)
    if nav:
        _h3(doc, "1b. Navamsa (D-9) Strength Layer (v15)")
        pl = ";  ".join(f"{k}: {v}" for k, v in nav.get("placements", {}).items())
        _para(doc, f"Key planet Navamsa placements (opening): {pl}.")
        for k in ("n3_text", "n5_text", "n6_text"):
            if nav.get(k):
                _para(doc, nav[k])
        if nav.get("moon_navamsa_change"):
            _para(doc, f"Moon navamsa changes at {nav['moon_navamsa_change']} — Mo: element flips in the timing table from there.", 9.5)

    # --- 1c. Nadi ---------------------------------------------------------
    nadi = getattr(ds, "nadi", None)
    if nadi:
        _h3(doc, "1c. SBC Nadi Classification (v15.1)")
        for line in nadi.get("text", "").split("\n"):
            _para(doc, line)

    # --- 2. Regime shifts -------------------------------------------------
    _h(doc, "2. Master Regime Shifts (Moon Sub-Lord, Rule A-11/A-12)")
    if not report.regime_shifts:
        _para(doc, "No Moon sub-lord change during market hours — the opening regime holds all session.")
    for rs in report.regime_shifts:
        _para(doc, (f"{rs.time_str()}  Moon sub change {rs.old_sublord}->{rs.new_sublord}.  "
                    f"Incoming sub-lord {rs.new_sublord} (house-set strength {rs.house_set_strength:+.1f}) "
                    f"-> post-shift regime: {rs.regime_label}."))

    # --- 3. Predictive path -----------------------------------------------
    _h(doc, "3. Predictive Intraday Path")
    _para(doc, ("Orange markers = Tier-1 / regime-shift timestamps. Path is directional shape, not price "
                "target. Windows are candidates until price confirms."), 9.5)
    png = _path_chart_png(report)
    if png:
        doc.add_picture(io.BytesIO(png), width=Pt(460))
    else:
        _para(doc, "(Chart available in the web report; matplotlib not installed for document export.)", 9)

    # --- 4. Timing table --------------------------------------------------
    _h(doc, "4. Trend Change Timing Table")
    rows = []
    for tc in report.trend_changes:
        rows.append([
            tc.time_str(), tc.grade.value, tc.direction.value, tc.trigger,
            tc.asc_sub_house_set, tc.hora_lord,
            getattr(tc, "navamsa_d9", ""), getattr(tc, "nadi", ""),
        ])
    _table(doc, ["Time", "Grade", "Direction", "Trigger (Rule A-8/A-11)",
                 "Asc sub house-set", "Hora", "Navamsa D-9", "Nadi"], rows)
    rev = [tc for tc in report.trend_changes if "REVERSAL" in tc.grade.value.upper()]
    if rev:
        items = "; ".join(f"{tc.time_str()} {tc.direction.value} (score {tc.score:+.2f}, {tc.notes or tc.trigger})"
                          for tc in rev)
        _para(doc, (f"REVERSAL CANDIDATE signals (v14): {items}.  These signals are generated when score "
                    "exhaustion in one direction coincides with pure-set alarms in the opposite direction — "
                    "a high-conviction reversal setup. Confirm with price: first 5 candles must move ≥ 15 pts "
                    "in the reversal direction."), 9.5)

    # --- 5. Trading view --------------------------------------------------
    _h(doc, "5. Intraday Trading View")
    sb, sbe = report.strongest_bullish, report.strongest_bearish
    parts = []
    if sb:
        parts.append(f"Strongest bullish candidate window: {sb.time_range_str()} (peak score {sb.score:.2f}).")
    if sbe:
        parts.append(f"Strongest bearish candidate window: {sbe.time_range_str()} (deepest score {sbe.score:.2f}).")
    if parts:
        _para(doc, "  ".join(parts))
    if report.exhaustion_warnings:
        _para(doc, "Score exhaustion warning: " + " | ".join(report.exhaustion_warnings), 9.5)
    _para(doc, ("Playbook: trade Tier-1 flips only, in the direction of the active Moon sub-lord regime. Longs "
                "are valid only in buy-on-dip regimes; shorts only in sell-on-rise regimes. EXCEPTION (v14): "
                "REVERSAL CANDIDATE signals may trade AGAINST the regime when exhaustion + pure-set alarms "
                "confirm — these are early reversal entries with elevated conviction. Every candidate window "
                "requires price confirmation: the first 5 five-minute candles must move at least 15 NIFTY "
                "points in the predicted direction, plus normal chart confirmation such as reversal candle, "
                "VWAP reclaim/loss, or CPR level break. Maximum three Tier-1 trades per session (Rule D-3)."))
    _para(doc, ("Bank Nifty calibration: widen timing tolerance to ±8 minutes, expect one-grade-stronger "
                "response to Moon-based triggers, and trade Rahu-sub windows at half size, breakout-only."))

    # --- 6. Disclaimer ----------------------------------------------------
    _h(doc, "6. Disclaimer")
    _para(doc, ("Educational research output combining KP astrology with systematic trading discipline. Not "
                "investment advice. Derivatives trading involves substantial risk of loss. Validate all signals "
                "with independent backtesting before deploying capital."), 9.5)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _path_chart_png(report):
    """Cumulative signal-score path with Tier-1 / regime markers; None without matplotlib."""
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
    except Exception:
        return None
    sig = sorted(report.trend_changes, key=lambda s: s.timestamp)
    if not sig:
        return None
    xs, ys, cum = [], [], 0.0
    marks = []
    for s in sig:
        step = s.score if s.direction.value == "Bullish" else -s.score if s.direction.value == "Bearish" else 0.0
        cum += step
        x = s.timestamp.hour + s.timestamp.minute / 60
        xs.append(x)
        ys.append(cum)
        if "TIER-1" in s.grade.value or "REGIME" in s.grade.value:
            marks.append((x, cum))
    fig, ax = plt.subplots(figsize=(7.2, 2.6), dpi=150)
    ax.plot(xs, ys, color="#2B4C7E", linewidth=1.6)
    ax.axhline(0, color="#999", linewidth=0.6)
    if marks:
        ax.scatter([m[0] for m in marks], [m[1] for m in marks], color="#E8842B", s=18, zorder=3)
    ax.set_xlim(9.25, 15.5)
    ax.set_xticks([9.25, 10, 11, 12, 13, 14, 15, 15.5])
    ax.set_xticklabels(["9:15", "10:00", "11:00", "12:00", "13:00", "14:00", "15:00", "15:30"], fontsize=7)
    ax.tick_params(axis="y", labelsize=7)
    ax.set_ylabel("cumulative score", fontsize=7)
    ax.grid(alpha=0.25)
    for sp in ("top", "right"):
        ax.spines[sp].set_visible(False)
    buf = io.BytesIO()
    fig.tight_layout()
    fig.savefig(buf, format="png")
    plt.close(fig)
    return buf.getvalue()


__all__ = ["render_docx"]
